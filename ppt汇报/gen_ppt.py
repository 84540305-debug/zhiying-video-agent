# -*- coding: utf-8 -*-
"""知影 ZhiYing 汇报 PPT 生成器（python-pptx）"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ============ 颜色 ============
GREEN = RGBColor(0x2E, 0x7D, 0x4F)
GREEN_L = RGBColor(0x5A, 0xA8, 0x7A)
GOLD = RGBColor(0xFF, 0xD9, 0x7A)
BG = RGBColor(0xF3, 0xF7, 0xF3)
DARK = RGBColor(0x1C, 0x32, 0x26)
MUTED = RGBColor(0x54, 0x70, 0x5E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF0, 0xF7, 0xF3)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

blank = prs.slide_layouts[6]

def add_bg(slide, color):
    bg = slide.background; fill = bg.fill; fill.solid(); fill.fore_color.rgb = color

def add_bar(slide, x=0, y=0, w=Emu(180000), h=None):
    if h is None: h = prs.slide_height
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    bar.fill.solid(); bar.fill.fore_color.rgb = GREEN; bar.line.fill.background()
    return bar

def add_text(slide, x, y, w, h, text, size=16, bold=False, color=DARK, align='left', font='微软雅黑'):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = {'left':1,'center':2,'right':3}[align]
    run = p.add_run(); run.text = text
    run.font.size = Pt(size); run.font.bold = bold
    run.font.name = font; run.font.color.rgb = color
    return tb

def add_pic(slide, path, x, y, w=None, h=None):
    """插入图片"""
    import os
    full = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets', path)
    if not os.path.exists(full): return None
    kw = {}
    if w: kw['width'] = w
    if h: kw['height'] = h
    return slide.shapes.add_picture(full, x, y, **kw)

def add_points(slide, x, y, w, h, lines, size=13, color=MUTED):
    """表格下方文字要点"""
    return add_multiline(slide, x, y, w, h, lines, size=size, color=color)

def add_multiline(slide, x, y, w, h, lines, size=14, color=DARK, font='微软雅黑'):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, (text, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        run = p.add_run(); run.text = text
        run.font.size = Pt(size); run.font.bold = bold
        run.font.name = font; run.font.color.rgb = color
    return tb

def add_code(slide, x, y, w, h, code, size=11):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(code.strip().split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0); p.line_spacing = 1.2
        run = p.add_run(); run.text = line
        run.font.size = Pt(size); run.font.name = 'Consolas'; run.font.color.rgb = GREEN
    # 背景色
    tb.fill.solid(); tb.fill.fore_color.rgb = LIGHT_BG
    tb.line.color.rgb = GREEN; tb.line.width = Pt(0.5)
    return tb

def add_table(slide, x, y, w, rows_data, col_widths=None, header_color=GREEN):
    n_rows = len(rows_data); n_cols = len(rows_data[0])
    total_h = Inches(0.4 * n_rows)
    table_shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, total_h)
    table = table_shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = cw
    for ri, row in enumerate(rows_data):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11 if ri > 0 else 12)
                p.font.bold = ri == 0
                p.font.name = '微软雅黑'
                p.font.color.rgb = WHITE if ri == 0 else DARK
            if ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = header_color
            elif ri % 2 == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = LIGHT_BG
    return table

def footer(slide, page_num, total=20):
    add_text(slide, Inches(0.5), Inches(7.0), Inches(4), Inches(0.4),
             '知影 ZhiYing', size=12, color=MUTED)
    add_text(slide, Inches(11), Inches(7.0), Inches(2), Inches(0.4),
             f'{page_num} / {total}', size=12, color=MUTED, align='right')

def title_bar(slide, title, page_num):
    """标准内容页：标题栏 + 页脚"""
    add_bg(slide, WHITE)
    add_bar(slide, 0, 0, Emu(120000))  # 左侧装饰条
    add_text(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8),
             title, size=28, bold=True, color=GREEN)
    # 标题下划线
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.15), Inches(0.8), Emu(40000))
    line.fill.solid(); line.fill.fore_color.rgb = GREEN; line.line.fill.background()
    footer(slide, page_num)

def section_page(page_num, title, subtitle=''):
    """章节扉页"""
    slide = prs.slides.add_slide(blank)
    add_bg(slide, GREEN)
    bar = add_bar(slide, 0, 0, Emu(200000))
    bar.fill.fore_color.rgb = GOLD
    add_text(slide, Inches(0.5), Inches(2.5), Inches(12.3), Inches(1.5),
             title, size=40, bold=True, color=WHITE, align='center')
    if subtitle:
        add_text(slide, Inches(0.5), Inches(4.0), Inches(12.3), Inches(0.8),
                 subtitle, size=18, color=GOLD, align='center')
    add_text(slide, Inches(0.5), Inches(5.5), Inches(12.3), Inches(0.5),
             '知影 ZhiYing · 大学生军事理论课 AI 视频成片 Agent', size=12, color=RGBColor(0xE0,0xF1,0xEA), align='center')
    return slide

def cover_page(title, subtitle, date_str):
    """封面"""
    slide = prs.slides.add_slide(blank)
    add_bg(slide, GREEN)
    bar = add_bar(slide, 0, 0, Emu(250000))
    bar.fill.fore_color.rgb = GOLD
    # 金星装饰
    star = slide.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(5.67), Inches(1.2), Inches(2), Inches(2))
    star.fill.solid(); star.fill.fore_color.rgb = GOLD; star.line.fill.background()
    add_text(slide, Inches(0.5), Inches(3.3), Inches(12.3), Inches(1.2),
             title, size=40, bold=True, color=WHITE, align='center')
    add_text(slide, Inches(0.5), Inches(4.5), Inches(12.3), Inches(0.8),
             subtitle, size=18, color=GOLD, align='center')
    add_text(slide, Inches(0.5), Inches(6.0), Inches(12.3), Inches(0.5),
             date_str, size=14, color=RGBColor(0xE0,0xF1,0xEA), align='center')
    return slide

# ============ 01 封面 ============
cover_page('知影 ZhiYing', '大学生军事理论课 AI 视频成片 Agent · 搭建汇报', '2026年9月10日')

# ============ 02 目录 ============
slide = prs.slides.add_slide(blank)
add_bg(slide, WHITE)
add_bar(slide, 0, 0, Emu(120000))
add_text(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8), '目 录', size=32, bold=True, color=GREEN)
line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.15), Inches(0.8), Emu(40000))
line.fill.solid(); line.fill.fore_color.rgb = GREEN; line.line.fill.background()
toc = [
    ('01', '项目概述', '核心能力 · 在线成果'),
    ('02', '技术架构', '技术栈 · 三栏布局'),
    ('03', '功能详解', '三模式 · 七步 · 知识库 · API'),
    ('04', '部署上线', 'GitHub · Pages · 代理'),
    ('05', '项目成果', '总结 · 在线地址'),
]
for i, (num, ch, desc) in enumerate(toc):
    y = Inches(1.6 + i * 1.1)
    # 编号圆
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), y, Inches(0.5), Inches(0.5))
    circle.fill.solid(); circle.fill.fore_color.rgb = GREEN; circle.line.fill.background()
    add_text(slide, Inches(0.8), y, Inches(0.5), Inches(0.5), num, size=14, bold=True, color=WHITE, align='center')
    add_text(slide, Inches(1.6), y, Inches(4), Inches(0.5), ch, size=18, bold=True, color=DARK)
    add_text(slide, Inches(6), y, Inches(6), Inches(0.5), desc, size=14, color=MUTED)
footer(slide, 2)

# ============ 03 第一章扉页 ============
section_page(3, '第一章 · 项目概述', 'Project Overview')

# ============ 04 核心能力总览 ============
slide = prs.slides.add_slide(blank)
title_bar(slide, '核心能力总览', 4)
add_table(slide, Inches(0.5), Inches(1.5), Inches(12.3), [
    ['能力', '说明'],
    ['三种创作模式', '图片轮播（文生图分镜）/ HTML 视频（网页动画）/ 演示页生成（PPT）'],
    ['七步流水线', '生成脚本 → 生成大纲 → 拆分分镜 → 生成画面 → 生成声音 → 预览修改 → 导出'],
    ['文生图 API', 'Pollinations.ai（免费免key）+ 可切换豆包/通义万相/Stability/DALL·E'],
    ['军事知识库', '8 大章节专业教学内容（国防/军事思想/信息化战争等）'],
    ['文件上传', '文稿注入脚本生成 / 图片作分镜素材 / 文档作参考资料'],
    ['双通道导出', 'python-pptx 脚本（.pptx）/ 可编辑 HTML（Gamma/Canva 兼容）'],
], col_widths=[Inches(2.5), Inches(9.8)])
add_points(slide, Inches(0.5), Inches(5.8), Inches(12.3), Inches(1.2), [
    ('要点：知影 ZhiYing 整合了从脚本生成到录屏导出的完整视频制作流程，让教学视频从"手动剪辑"升级为"AI 一键成片"。', False),
    ('三大模式覆盖不同教学场景：图片轮播适合知识科普、HTML 视频适合动态演示、演示页生成适合课堂讲解。', False),
    ('内置军事理论课专业知识库，输入军事关键词即可自动匹配专业教学内容，无需手动编写文案。', False),
], size=12, color=MUTED)

# ============ 05 在线成果与文件清单 ============
slide = prs.slides.add_slide(blank)
title_bar(slide, '在线成果与文件清单', 5)
# 插入首页截图
add_pic(slide, 'homepage.png', Inches(0.5), Inches(1.5), w=Inches(7.5))
# 右栏：在线地址
add_text(slide, Inches(8.3), Inches(1.5), Inches(4.5), Inches(0.4), '在线访问', size=16, bold=True, color=GREEN)
add_multiline(slide, Inches(8.3), Inches(2.0), Inches(4.5), Inches(3), [
    ('🌐 在线网站', True),
    ('84540305-debug.github.io', False),
    ('  /zhiying-video-agent/', False),
    ('', False),
    ('📦 代码仓库', True),
    ('github.com/', False),
    ('  84540305-debug/', False),
    ('  zhiying-video-agent', False),
], size=12, color=DARK)
add_points(slide, Inches(0.5), Inches(6.0), Inches(12.3), Inches(1), [
    ('要点：网站已上线 GitHub Pages，任何人打开链接即可体验全部功能，无需安装。push 代码后自动重建。', False),
], size=12, color=MUTED)

# ============ 06 第二章扉页 ============
section_page(6, '第二章 · 技术架构', 'Technical Architecture')

# ============ 07 技术栈八层 ============
slide = prs.slides.add_slide(blank)
title_bar(slide, '技术栈八层', 7)
add_table(slide, Inches(0.5), Inches(1.5), Inches(12.3), [
    ['层', '技术', '说明'],
    ['前端', '原生 HTML+CSS+JS', '单文件，零依赖，零构建'],
    ['文生图', 'Pollinations.ai API', '免费免key，渐进式加载'],
    ['语音合成', 'Web Speech API', '浏览器内置中文 TTS'],
    ['背景音乐', 'WebAudio API', '程序化生成 3 种风格'],
    ['录屏导出', 'getDisplayMedia+MediaRecorder', '屏幕捕获→.webm 下载'],
    ['持久化', 'localStorage', '历史项目存储'],
    ['版本控制', 'Git + GitHub', 'main 分支'],
    ['在线托管', 'GitHub Pages', 'push 后自动重建'],
], col_widths=[Inches(1.5), Inches(4.5), Inches(6.3)])
add_points(slide, Inches(0.5), Inches(5.8), Inches(12.3), Inches(1.2), [
    ('要点：采用纯前端 B/S 架构，单文件零依赖，所有功能在浏览器内完成，无需后端服务器。', False),
    ('文生图通过 Pollinations.ai 免费 API 实现，无需 API Key；可一键切换豆包/Stability/DALL·E 等付费模型。', False),
    ('录屏导出利用浏览器原生 getDisplayMedia API，无需安装录屏软件，导出 .webm 视频文件。', False),
], size=12, color=MUTED)

# ============ 08 三栏布局架构 ============
slide = prs.slides.add_slide(blank)
title_bar(slide, '三栏布局架构', 8)
# 插入编辑器截图
add_pic(slide, 'editor.png', Inches(0.5), Inches(1.5), w=Inches(12.3))
add_points(slide, Inches(0.5), Inches(6.0), Inches(12.3), Inches(1), [
    ('要点：首页 → 点击模式卡片 → 进入创作页。左栏历史项目（封面缩略图），中栏预览+工具栏+分镜列表，右栏 AI 对话。', False),
    ('三栏协同：左侧管理项目、中间预览编辑、右侧 AI 创作对话，全程不切换页面。', False),
], size=12, color=MUTED)

# ============ 09 第三章扉页 ============
section_page(9, '第三章 · 功能详解', 'Feature Details')

# ============ 10 三种创作模式 ============
slide = prs.slides.add_slide(blank)
title_bar(slide, '三种创作模式', 10)
add_table(slide, Inches(0.5), Inches(1.5), Inches(12.3), [
    ['模式', '说明', '入口函数'],
    ['图片轮播', '文生图 AI 生成静态图片，Ken Burns 镜头动效轮播', "enterEditor('image')"],
    ['HTML 视频', 'AI 生成网页动画（极光/粒子/轨道/数据律动/打字机）', "enterEditor('html')"],
    ['演示页生成', 'AI 生成 5-8 页演示稿，导出 .pptx 或可编辑 HTML', "enterEditor('ppt')"],
], col_widths=[Inches(1.5), Inches(8.8), Inches(2.0)])
add_code(slide, Inches(0.5), Inches(4.5), Inches(12.3), Inches(2.0), """// 模式切换
function enterEditor(mode){
  state.mode = mode;
  document.body.classList.toggle('mode-ppt', mode === 'ppt');
  // mode='image' → 图片轮播 | 'html' → HTML动画 | 'ppt' → 演示页
}""", size=12)
add_points(slide, Inches(0.5), Inches(6.3), Inches(12.3), Inches(0.8), [
    ('要点：三种模式共享同一套七步流水线，差异仅在画面生成方式。模式切换只需一行代码（mode-ppt class）。', False),
], size=12, color=MUTED)

# ============ 11 七步流水线 ============
slide = prs.slides.add_slide(blank)
title_bar(slide, '七步创作流水线', 11)
# 插入七步流程图
add_pic(slide, 'pipeline.png', Inches(0.5), Inches(1.5), w=Inches(12.3))
add_points(slide, Inches(0.5), Inches(5.5), Inches(12.3), Inches(1.5), [
    ('要点：用户输入提示词后，对话区顶部生成「创作流水线·7步」追踪器，随进度实时高亮"进行中/已完成"。', False),
    ('三级内容路由：上传文稿 > 军事知识库 > 通用模板。命中军事关键词时自动使用专业教学内容。', False),
    ('每步完成后对话区产出对应内容（脚本全文/大纲表格/分镜列表/画面进度/声音说明），用户可实时查看。', False),
], size=12, color=MUTED)

# ============ 12 军事知识库 ============
slide = prs.slides.add_slide(blank)
title_bar(slide, '军事理论课知识库', 12)
# 插入知识库配图
add_pic(slide, 'knowledge_base.png', Inches(0.5), Inches(1.5), w=Inches(12.3))
add_points(slide, Inches(0.5), Inches(5.8), Inches(12.3), Inches(1.2), [
    ('要点：matchMilitaryTopic(prompt) 关键词匹配 → 命中时脚本/演示稿自动使用专业教学旁白文案。', False),
    ('四套军事专属配色自动匹配：国防→军绿迷彩、战争→战旗红金、信息化→夜战深蓝、高技术→钢铁灰蓝。', False),
    ('非军事主题兜底也加军事元素修饰词，确保全站视觉统一性（military themed, soldiers, national defense）。', False),
], size=12, color=MUTED)

# ============ 13 文生图 API ============
slide = prs.slides.add_slide(blank)
title_bar(slide, '文生图 API 接入', 13)
add_table(slide, Inches(0.5), Inches(1.5), Inches(12.3), [
    ['模型', '函数', '费用', 'CORS'],
    ['Pollinations.ai', '_callPollinations()', '免费', '✅ 直调'],
    ['豆包/火山引擎', '_callDoubao()', '按量', '❌ 需代理'],
    ['通义万相', '_callWanx()', '按量', '❌ 需代理'],
    ['Stability AI', '_callStability()', '按量', '✅ 直调'],
    ['OpenAI DALL·E 3', '_callDallE()', '按量', '✅ 直调'],
], col_widths=[Inches(2.5), Inches(3.5), Inches(1.5), Inches(2.0)])
add_code(slide, Inches(0.5), Inches(5.0), Inches(12.3), Inches(1.8), """// 路由函数：根据配置选择模型
function _callImageAPI(prompt, w, h){
  switch(IMG_CONFIG.model){
    case 'doubao':    return _callDoubao(prompt, w, h);
    case 'stability': return _callStability(prompt, w, h);
    default:          return _callPollinations(prompt, w, h);
  }
}""", size=11)
add_points(slide, Inches(0.5), Inches(6.5), Inches(12.3), Inches(0.6), [
    ('要点：首页 ⚙️ 设置弹窗可一键切换模型，API Key 仅存本地 localStorage。渐进式加载（Canvas占位→3秒后真实图淡入）。', False),
], size=11, color=MUTED)

# ============ 14 演示页生成与导出 ============
slide = prs.slides.add_slide(blank)
title_bar(slide, '演示页生成与导出', 14)
# 插入幻灯片截图
add_pic(slide, 'slide.png', Inches(0.5), Inches(1.4), w=Inches(7.5))
# 右栏
add_text(slide, Inches(8.3), Inches(1.5), Inches(4.5), Inches(0.4), '双通道导出', size=16, bold=True, color=GREEN)
add_table(slide, Inches(8.3), Inches(2.0), Inches(4.5), [
    ['导出方式', '产物'],
    ['PPTX', '.py脚本→.pptx'],
    ['HTML', '单文件→Gamma'],
], col_widths=[Inches(1.8), Inches(2.7)])
add_points(slide, Inches(0.5), Inches(5.5), Inches(12.3), Inches(1.5), [
    ('要点：演示页模式复用七步流水线，按教学章节结构生成 5-8 页：封面→引入→核心概念→关键特征→应用→案例→总结。', False),
    ('四种版式自动选择：title 居中标题页 / content 项目符号页 / two-col 图文双栏 / summary 结语总结。', False),
    ('PPTX 导出生成 python-pptx 脚本（含全部页面数据），本地 pip install 后运行即产出标准 .pptx 文件。', False),
], size=12, color=MUTED)

# ============ 15 第四章扉页 ============
section_page(15, '第四章 · 部署上线', 'Deployment & Online')

# ============ 16 GitHub 仓库部署 ============
slide = prs.slides.add_slide(blank)
title_bar(slide, 'GitHub 仓库部署', 16)
add_text(slide, Inches(0.5), Inches(1.5), Inches(5.5), Inches(0.4), '部署步骤', size=18, bold=True, color=GREEN)
add_multiline(slide, Inches(0.5), Inches(2.0), Inches(5.5), Inches(4), [
    ('1. 安装 Git + 注册 GitHub', False),
    ('2. 生成 PAT（classic, repo scope）', False),
    ('3. git init + commit', False),
    ('4. 创建远程仓库', False),
    ('5. git push 推送代码', False),
], size=14, color=DARK)
add_text(slide, Inches(0.5), Inches(5.5), Inches(5.5), Inches(0.4), '认证方式', size=14, bold=True, color=GREEN)
add_text(slide, Inches(0.5), Inches(5.9), Inches(5.5), Inches(0.8),
         'Classic PAT（ghp_ 开头）\nrepo scope，用完即删', size=12, color=MUTED)
# 右栏：代码
add_code(slide, Inches(6.5), Inches(1.5), Inches(6.3), Inches(5.2), """# 本地初始化
git init -b main
git config user.name "用户名"
git config user.email "邮箱"
git add -A
git commit -m "初始版本"

# 推送（PAT 嵌入 URL，不持久化）
export GIT_TERMINAL_PROMPT=0
git -c credential.helper= push \\
  "https://用户名:ghp_xxx@github.com/
   用户名/zhiying-video-agent.git" \\
  main:main""", size=11)
add_points(slide, Inches(0.5), Inches(6.3), Inches(12.3), Inches(0.8), [
    ('要点：Classic PAT（ghp_ 开头）勾选 repo scope 即可，用完即删。GIT_TERMINAL_PROMPT=0 + credential.helper= 避免静默挂起。', False),
], size=11, color=MUTED)

# ============ 17 GitHub Pages ============
slide = prs.slides.add_slide(blank)
title_bar(slide, 'GitHub Pages 在线部署', 17)
# 插入部署流程图
add_pic(slide, 'pages_deploy.png', Inches(0.5), Inches(1.5), w=Inches(12.3))
add_points(slide, Inches(0.5), Inches(5.0), Inches(12.3), Inches(2), [
    ('要点：用 PAT 调 GitHub API（POST /repos/.../pages）启用 Pages，source 设为 main 分支根目录。', False),
    ('首次构建约 30-60 秒，之后每次 git push 后 Pages 自动重建，在线网站自动同步更新。', False),
    ('任何人打开 https://84540305-debug.github.io/zhiying-video-agent/ 即可直接体验全部功能。', False),
], size=13, color=MUTED)

# ============ 18 CORS 代理 ============
slide = prs.slides.add_slide(blank)
title_bar(slide, 'CORS 代理部署', 18)
# 左栏：Vercel
add_text(slide, Inches(0.5), Inches(1.5), Inches(5.8), Inches(0.4), '方案 A：Vercel Edge（推荐）', size=16, bold=True, color=GREEN)
add_multiline(slide, Inches(0.5), Inches(2.0), Inches(5.8), Inches(3), [
    ('1. vercel.com/new → GitHub 登录', False),
    ('2. Import 仓库', False),
    ('3. Deploy（30秒）', False),
    ('4. 代理地址：', False),
    ('  vercel.app/api/proxy/doubao', False),
], size=13, color=DARK)
# 右栏：路由表
add_text(slide, Inches(7), Inches(1.5), Inches(5.8), Inches(0.4), '路由映射', size=16, bold=True, color=GREEN)
add_table(slide, Inches(7), Inches(2.0), Inches(5.8), [
    ['路径', '转发到'],
    ['/doubao/*', 'ark.cn-beijing.volces.com'],
    ['/wanx/*', 'dashscope.aliyuncs.com'],
    ['/openai/*', 'api.openai.com'],
    ['/stability/*', 'api.stability.ai'],
], col_widths=[Inches(2.5), Inches(3.3)])
add_text(slide, Inches(0.5), Inches(5.5), Inches(12.3), Inches(0.4), '方案 B：Cloudflare Workers', size=16, bold=True, color=GREEN)
add_text(slide, Inches(0.5), Inches(6.0), Inches(12.3), Inches(0.6),
         'workers.cloudflare.com → Create Worker → 粘贴 worker.js → Deploy → 得到 xxx.workers.dev', size=13, color=MUTED)

# ============ 19 第五章扉页 ============
section_page(19, '第五章 · 项目成果', 'Project Outcomes')

# ============ 20 结束页 ============
slide = prs.slides.add_slide(blank)
add_bg(slide, GREEN)
bar = add_bar(slide, 0, 0, Emu(250000))
bar.fill.fore_color.rgb = GOLD
star = slide.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(5.67), Inches(1.0), Inches(2), Inches(2))
star.fill.solid(); star.fill.fore_color.rgb = GOLD; star.line.fill.background()
add_text(slide, Inches(0.5), Inches(3.0), Inches(12.3), Inches(1.0),
         '感谢聆听', size=40, bold=True, color=WHITE, align='center')
add_text(slide, Inches(0.5), Inches(4.2), Inches(12.3), Inches(0.5),
         '🌐 https://84540305-debug.github.io/zhiying-video-agent/', size=16, color=GOLD, align='center')
add_text(slide, Inches(0.5), Inches(4.8), Inches(12.3), Inches(0.5),
         '📦 https://github.com/84540305-debug/zhiying-video-agent', size=16, color=GOLD, align='center')
add_text(slide, Inches(0.5), Inches(5.8), Inches(12.3), Inches(0.5),
         '知影 ZhiYing · 大学生军事理论课 AI 视频成片 Agent · 2026年9月', size=14, color=RGBColor(0xE0,0xF1,0xEA), align='center')

# ============ 保存 ============
output = os.path.join(os.path.dirname(os.path.abspath(__file__)), '知影ZhiYing汇报.pptx')
prs.save(output)
print(f'PPT已生成：{output}')
print(f'文件大小：{os.path.getsize(output)} bytes')
print(f'总页数：{len(prs.slides._sldIdLst)} 页')
